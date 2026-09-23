from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import BinaryIO, Iterable, Optional, Tuple, Union
from xml.etree import ElementTree as ET

from .schemas import Clause, Document, Side


PathLike = Union[str, Path]
WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def parse_document(
    source: Union[PathLike, bytes, BinaryIO],
    *,
    document_id: Optional[str] = None,
    document_label: Optional[str] = None,
) -> Document:
    data, name = _read_source(source)
    label = document_label or name
    doc_id = document_id or _slug(Path(name).stem or label)
    suffix = Path(name).suffix.lower()

    if suffix == ".docx":
        rows = _parse_docx(data)
    elif suffix == ".pdf":
        rows = _parse_pdf(data)
    elif suffix in {".xlsx", ".xlsm"}:
        rows = _parse_xlsx(data)
    elif suffix in {".txt", ".md", ""}:
        rows = _parse_text(data)
    else:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    clauses = _rows_to_clauses(rows)
    if not clauses:
        raise ValueError(f"No text extracted from {label}")
    return Document(id=doc_id, name=label, clauses=clauses)


def parse_many(files: Iterable[Union[PathLike, bytes, BinaryIO]], *, side: Side) -> list[Document]:
    documents: list[Document] = []
    for index, file_obj in enumerate(files, start=1):
        documents.append(parse_document(file_obj, document_id=f"{side}{index}"))
    return documents


def document_clause_index(before: list[Document], after: list[Document]) -> dict[tuple[str, str, str], tuple[Document, Clause]]:
    index: dict[tuple[str, str, str], tuple[Document, Clause]] = {}
    for side, documents in (("before", before), ("after", after)):
        for document in documents:
            for clause in document.clauses:
                index[(side, document.id, clause.id)] = (document, clause)
    return index


def _read_source(source: Union[PathLike, bytes, BinaryIO]) -> Tuple[bytes, str]:
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_bytes(), path.name

    if isinstance(source, bytes):
        return source, "document.txt"

    name = getattr(source, "name", "upload.bin")
    if hasattr(source, "getvalue"):
        return source.getvalue(), name
    return source.read(), name


def _parse_docx(data: bytes) -> list[tuple[str, str, dict]]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml")

    root = ET.fromstring(xml)
    body = root.find(f"{WORD_NS}body")
    if body is None:
        return []

    rows: list[tuple[str, str, dict]] = []
    for node in body:
        tag = _strip_ns(node.tag)
        if tag == "p":
            text = _node_text(node)
            if text.strip():
                rows.append(("paragraph", text, {}))
        elif tag == "tbl":
            for table_row in node.iter(f"{WORD_NS}tr"):
                cells = []
                for cell in table_row.findall(f"{WORD_NS}tc"):
                    cell_text = _clean_text(_node_text(cell))
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    rows.append(("table_row", " | ".join(cells), {}))
    return rows


def _node_text(node: ET.Element) -> str:
    parts = []
    for child in node.iter():
        tag = _strip_ns(child.tag)
        if tag == "t" and child.text:
            parts.append(child.text)
        elif tag == "tab":
            parts.append(" ")
        elif tag in {"br", "cr"}:
            parts.append("\n")
    return "".join(parts)


def _parse_pdf(data: bytes) -> list[tuple[str, str, dict]]:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError as exc:
            raise ImportError("Install pypdf or PyPDF2 to parse PDF files.") from exc

    reader = PdfReader(io.BytesIO(data))
    rows = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            rows.append(("page", text, {"page": page_number}))
    return rows


def _parse_xlsx(data: bytes) -> list[tuple[str, str, dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("Install openpyxl to parse XLSX files.") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows: list[tuple[str, str, dict]] = []
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(("sheet_row", " | ".join(values), {"sheet": sheet.title, "row": row_number}))
    return rows


def _parse_text(data: bytes) -> list[tuple[str, str, dict]]:
    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")

    rows = []
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if paragraph:
            rows.append(("text", paragraph, {}))
    return rows


def _rows_to_clauses(rows: list[tuple[str, str, dict]]) -> list[Clause]:
    clauses: list[Clause] = []
    seen: dict[str, int] = {}
    counters = {"paragraph": 0, "table_row": 0, "sheet_row": 0, "page": 0, "text": 0}
    markers = {"paragraph": "p", "table_row": "table", "sheet_row": "sheet", "page": "page", "text": "txt"}

    for kind, text, meta in rows:
        text = _clean_text(text)
        if not text:
            continue
        for part in _split_long_text(text):
            clause_id = _clause_id(part)
            if not clause_id:
                counters[kind] += 1
                if kind == "sheet_row":
                    clause_id = f"{markers[kind]}-{meta.get('sheet', 'sheet')}-{meta.get('row', counters[kind])}"
                elif kind == "page":
                    clause_id = f"page-{meta.get('page', counters[kind])}"
                else:
                    clause_id = f"{markers[kind]}-{counters[kind]}"

            seen[clause_id] = seen.get(clause_id, 0) + 1
            unique_id = clause_id if seen[clause_id] == 1 else f"{clause_id}-{seen[clause_id]}"
            clauses.append(Clause(id=unique_id, text=part))
    return clauses


def _clause_id(text: str) -> Optional[str]:
    match = re.match(r"^(\d+(?:\.\d+)+)\.?\s+", text)
    if match:
        return match.group(1)
    section = re.match(r"^(\d+)\.\s+[\wА-Яа-я]", text)
    if section:
        return section.group(1)
    return None


def _split_long_text(text: str, limit: int = 1800) -> list[str]:
    if len(text) <= limit:
        return [text]

    parts = []
    current = []
    current_len = 0
    for sentence in re.split(r"(?<=[.!?;:])\s+", text):
        if current and current_len + len(sentence) > limit:
            parts.append(" ".join(current).strip())
            current = [sentence]
            current_len = len(sentence)
        else:
            current.append(sentence)
            current_len += len(sentence) + 1
    if current:
        parts.append(" ".join(current).strip())
    return [part for part in parts if part]


def _clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _strip_ns(tag: str) -> str:
    return tag.split("}", 1)[-1]


def _slug(value: str) -> str:
    translit = {
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "е": "e",
        "ё": "e",
        "ж": "zh",
        "з": "z",
        "и": "i",
        "й": "y",
        "к": "k",
        "л": "l",
        "м": "m",
        "н": "n",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "у": "u",
        "ф": "f",
        "х": "h",
        "ц": "c",
        "ч": "ch",
        "ш": "sh",
        "щ": "sch",
        "ы": "y",
        "э": "e",
        "ю": "yu",
        "я": "ya",
    }
    lowered = value.lower()
    normalized = "".join(translit.get(ch, ch) for ch in lowered)
    normalized = re.sub(r"[^a-z0-9]+", "", normalized)
    return normalized[:32] or "doc"

